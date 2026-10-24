from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "FYP_II_Comprehensive_Plan_AR.docx"

NAVY = "16324F"
BLUE = "2E74B5"
TEAL = "1B7F79"
GOLD = "B7791F"
INK = "1F2937"
MUTED = "5F6B76"
LIGHT_BLUE = "E8EEF5"
LIGHT_TEAL = "E8F4F2"
LIGHT_GOLD = "FFF7E6"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
RED = "9B1C1C"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = align
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_run(run, size=11, bold=False, color=INK, font="Arial", italic=False):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rtl = run._element.get_or_add_rPr().find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        run._element.get_or_add_rPr().append(rtl)
    rtl.set(qn("w:val"), "1")


def add_page_number(paragraph):
    set_rtl(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    run = paragraph.add_run("صفحة ")
    set_run(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run2 = paragraph.add_run()
    set_run(run2, size=9, color=MUTED)
    run2._r.extend([begin, instr, separate, text, end])


def add_paragraph(doc, text="", *, size=11, bold=False, color=INK, after=6,
                  before=0, align=WD_ALIGN_PARAGRAPH.RIGHT, italic=False,
                  keep_with_next=False):
    p = doc.add_paragraph()
    set_rtl(p, align)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_with_next = keep_with_next
    if text:
        r = p.add_run(text)
        set_run(r, size=size, bold=bold, color=color, italic=italic)
    return p


def add_bullet(doc, text, level=0, color=INK):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_rtl(p)
    p.paragraph_format.left_indent = Inches(0)
    p.paragraph_format.right_indent = Inches(0.375 + 0.25 * level)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run(r, size=11, color=color)
    return p


def add_numbered(doc, text, level=0):
    p = doc.add_paragraph(style="List Number" if level == 0 else "List Number 2")
    set_rtl(p)
    p.paragraph_format.right_indent = Inches(0.375 + 0.25 * level)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run(r, size=11)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_rtl(p)
    r = p.add_run(text)
    sizes = {1: 16, 2: 13, 3: 12}
    colors = {1: BLUE, 2: BLUE, 3: NAVY}
    set_run(r, size=sizes[level], bold=True, color=colors[level])
    return p


def add_callout(doc, title, text, fill=LIGHT_TEAL, accent=TEAL):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    set_rtl(p)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run(r, size=11, bold=True, color=accent)
    p2 = cell.add_paragraph()
    set_rtl(p2)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.25
    r2 = p2.add_run(text)
    set_run(r2, size=10.5)
    add_paragraph(doc, "", after=3)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        set_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_run(r, size=font_size, bold=True, color=NAVY)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            if row_idx % 2 == 1:
                set_cell_shading(cells[idx], "FAFBFC")
            p = cells[idx].paragraphs[0]
            set_rtl(p)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(value))
            set_run(r, size=font_size)
    set_table_geometry(table, widths)
    add_paragraph(doc, "", after=3)
    return table


def page_break(doc):
    doc.add_page_break()


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level, size, before, after, color in [
        (1, 16, 18, 10, BLUE),
        (2, 13, 14, 7, BLUE),
        (3, 12, 10, 5, NAVY),
    ]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ["List Bullet", "List Bullet 2", "List Number", "List Number 2"]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
        style.font.size = Pt(11)


def build_document():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    set_rtl(header)
    header.paragraph_format.space_after = Pt(0)
    hr = header.add_run("خطة مشروع التخرج 2  |  منصة تحليلات JustPlay")
    set_run(hr, size=8.5, bold=True, color=MUTED)
    add_page_number(section.footer.paragraphs[0])

    # Editorial cover
    add_paragraph(doc, "مشروع التخرج الثاني", size=12, bold=True, color=GOLD,
                  after=20, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "الخطة التنفيذية الشاملة", size=28, bold=True, color=NAVY,
                  after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "منصة تفاعلية لتحليل سلوك المستخدمين والربحية وتجارب A/B في تطبيق JustPlay",
                  size=15, bold=True, color=TEAL, after=34,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "امتداد عملي لمشروع التخرج الأول، من التحليل الاستكشافي على عينة إلى نظام تحليلي كامل مبني على جميع البيانات",
                  size=11.5, color=MUTED, after=42,
                  align=WD_ALIGN_PARAGRAPH.CENTER)

    cover_rows = [
        ("الطالب", "محمد إبراهيم روقة"),
        ("الرقم الجامعي", "20221013"),
        ("المشرف", "أ.د. إيهاب زقوت"),
        ("القسم", "هندسة البرمجيات"),
        ("الوثيقة", "خطة تنفيذ وإدارة وتسليم FYP II"),
        ("تاريخ إعداد الخطة", "13 أكتوبر 2026"),
        ("الإصدار", "1.0"),
    ]
    add_table(doc, ["البند", "التفاصيل"], cover_rows, [2700, 6660], header_fill=LIGHT_GOLD, font_size=10)
    add_callout(doc, "القرار المركزي", "تعتمد النتائج النهائية على كامل بيانات المستخدمين والمعاملات. تبقى البيانات الخام داخل مجلد المشروع محلياً، بينما يعتمد الداشبورد على ملفات نتائج مجمعة صغيرة وآمنة.", LIGHT_TEAL, TEAL)

    page_break(doc)
    add_heading(doc, "1. الملخص التنفيذي", 1)
    add_paragraph(doc, "يحوّل مشروع التخرج الثاني العمل المنجز في FYP I من تحليل أولي داخل Jupyter Notebook إلى منتج تحليلي متكامل وقابل للتشغيل والعرض والنشر. سيعالج النظام مجموعة البيانات الكاملة، ويطبق تعريفات صحيحة للمؤشرات، وينفذ اختبارات إحصائية موثقة، ثم يعرض النتائج في داشبورد Streamlit تفاعلي سريع.")
    add_callout(doc, "قاعدة العمل", "لا تُستخدم عينة الـ50,000 مستخدم لإصدار الحكم النهائي. تبقى نتائجها موثقة كمرحلة استكشافية في FYP I، بينما تُستخرج نتائج FYP II من كامل البيانات.", LIGHT_BLUE, BLUE)

    add_heading(doc, "2. الانتقال من FYP I إلى FYP II", 1)
    add_table(doc, ["البعد", "FYP I", "FYP II"], [
        ("البيانات", "عينة عشوائية ثابتة من 50,000 مستخدم", "كامل البيانات: 934,270 مستخدماً وجميع المعاملات"),
        ("الهدف", "إثبات الفكرة والتحليل الاستكشافي", "تحليل نهائي ونظام عملي قابل للاستخدام"),
        ("المخرجات", "Notebook ورسوم وتقرير أولي", "Pipeline + اختبارات + Dashboard + نشر + تقرير نهائي"),
        ("النتائج", "وصفية وأولية", "مدققة إحصائياً وموثقة وقابلة لإعادة الإنتاج"),
        ("الأداء", "قراءة CSV للعينة", "DuckDB/Parquet وملفات مجمعة محسنة"),
    ], [1900, 3300, 4160], font_size=9.2)

    add_heading(doc, "3. أهداف المشروع", 1)
    for item in [
        "بناء خط معالجة بيانات كامل وقابل لإعادة التشغيل على البيانات الأصلية دون تعديلها.",
        "تعريف مؤشرات الربحية والاحتفاظ والتفعيل بصورة دقيقة وقابلة للدفاع الأكاديمي.",
        "تقييم الـVariations الخمسة وصفياً وإحصائياً مع فواصل ثقة وأحجام أثر.",
        "تحليل قنوات الاستحواذ والدول وأنواع الإعلانات ومجموعات التثبيت الزمنية.",
        "إنشاء داشبورد تفاعلي يدعم الفلاتر، المقارنة، الاستكشاف، وتنزيل النتائج.",
        "توفير اختبارات آلية وتوثيق وتشغيل محلي ونشر آمن لنسخة العرض.",
        "إنتاج تقرير نهائي وعرض مناقشة يعكسان التطور الحقيقي بين FYP I وFYP II.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "4. معايير النجاح", 1)
    add_table(doc, ["المعيار", "شرط القبول"], [
        ("صحة البيانات", "تطابق المجاميع الأساسية مع الملفات الخام وتوثيق جميع الاستبعادات والتحويلات."),
        ("قابلية الإعادة", "تشغيل خط المعالجة من البداية بأمر واحد وإنتاج النتائج نفسها."),
        ("الدقة المنهجية", "تطبيق أهلية نوافذ D1/D7/D30 ومنع تحيز المستخدمين غير مكتملين المتابعة."),
        ("الاستدلال", "عدم إعلان فائز إلا مع دلالة إحصائية وأثر عملي وفاصل ثقة مناسب."),
        ("الأداء", "تحميل الصفحة الأساسية خلال ثوانٍ باستخدام ملفات مجمعة بدلاً من CSV الخام."),
        ("الخصوصية", "عدم نشر user_id أو المعاملات الفردية أو ملفات البيانات الخام."),
        ("جودة المنتج", "نجاح الاختبارات، وضوح الواجهات، وإمكانية تشغيل نسخة محلية احتياطية."),
    ], [2500, 6860], font_size=9.3)

    page_break(doc)
    add_heading(doc, "5. نطاق المشروع", 1)
    add_heading(doc, "5.1 داخل النطاق", 2)
    for item in [
        "البيانات الكاملة للمستخدمين والمعاملات المتاحة في المشروع.",
        "تحليل Variations، الدول، القنوات، الإيرادات، السحوبات، وأنواع الإعلانات.",
        "تحليل Cohorts ونشاط D1/D3/D7/D14/D30 وفق أهلية زمنية واضحة.",
        "اختبارات إحصائية، فواصل ثقة، أحجام أثر، وتصحيح تعدد المقارنات.",
        "تطبيق Streamlit متعدد الصفحات مع Plotly وملفات Parquet/DuckDB.",
        "تشغيل محلي، نسخة عرض سحابية، توثيق، اختبارات، وتقرير نهائي.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "5.2 خارج النطاق حالياً", 2)
    for item in [
        "تعديل تطبيق JustPlay الإنتاجي أو التكامل المباشر مع أنظمة الشركة.",
        "المعالجة اللحظية للبيانات؛ النظام تحليلي Batch حسب البيانات المتوفرة.",
        "حساب ROAS الحقيقي ما لم تتوفر تكلفة الإنفاق الإعلاني لكل قناة أو حملة.",
        "الاحتفاظ المبني على فتح التطبيق ما لم تتوفر بيانات Sessions أو Login events.",
        "التنبؤ بـLTV أو Churn ضمن النسخة الأساسية؛ يمكن إضافتهما كأهداف توسعية.",
    ]:
        add_bullet(doc, item, color=MUTED)

    add_heading(doc, "6. سياسة البيانات وحماية النسخة الأصلية", 1)
    add_numbered(doc, "تبقى ملفات مشروع التخرج الأول الأصلية دون تعديل أو إعادة تسمية.")
    add_numbered(doc, "تُستخدم نسخ مستقلة داخل FYP_II/data/raw للقراءة فقط أثناء التحليل.")
    add_numbered(doc, "لا تُكتب نتائج أو ملفات مؤقتة داخل data/raw؛ جميع المخرجات تذهب إلى processed أو marts.")
    add_numbered(doc, "لا تُرفع الملفات الخام إلى GitHub أو Streamlit Cloud.")
    add_numbered(doc, "تُقارن أحجام الملفات وقيم SHA-256 للتأكد من سلامة النسخ.")
    add_numbered(doc, "تزال المعرفات الفردية من جميع ملفات الداشبورد والتقارير القابلة للنشر.")

    add_heading(doc, "7. بنية المجلد المعتمدة", 1)
    folder_rows = [
        ("docs", "الخطة، قاموس البيانات، تعريفات المؤشرات، ومحاضر القرارات."),
        ("data/raw", "نسخة غير معدلة من ملف المستخدمين وملف الإيرادات الكاملين."),
        ("data/processed", "بيانات منظفة ومهيأة داخلياً؛ لا تُنشر."),
        ("data/marts", "ملفات النتائج المجمعة التي يقرأها الداشبورد."),
        ("src/pipeline", "تحميل البيانات، التدقيق، التنظيف، الربط، والتحويل إلى Parquet."),
        ("src/analytics", "حساب KPIs، Cohorts، الاختبارات، وقواعد التوصيات."),
        ("app/pages", "صفحات Streamlit ومكونات العرض والفلاتر."),
        ("tests", "اختبارات الوحدات والتكامل وجودة البيانات."),
        ("notebooks", "دفاتر الاستكشاف فقط؛ المنطق النهائي ينتقل إلى src."),
        ("reports/figures", "الرسوم النهائية المخصصة للتقرير والمناقشة."),
        ("config", "إعدادات المسارات، تعريف المؤشرات، وحدود أهلية Cohorts."),
        ("deployment", "Docker وملفات ومتطلبات النشر."),
        ("logs", "سجلات تشغيل Pipeline وتقارير التحقق."),
    ]
    add_table(doc, ["المجلد", "الغرض"], folder_rows, [2600, 6760], font_size=9.2)

    page_break(doc)
    add_heading(doc, "8. المعمارية التقنية", 1)
    add_paragraph(doc, "يتبع النظام مبدأ الفصل بين المعالجة الثقيلة والعرض الخفيف. تُقرأ البيانات الخام محلياً، ثم تُنظف وتُجمع وتُحفظ في طبقة تحليلية صغيرة. يقرأ الداشبورد طبقة النتائج فقط.")
    add_table(doc, ["الطبقة", "الأداة المقترحة", "المسؤولية"], [
        ("Raw Data", "CSV", "المصدر الأصلي غير المعدل."),
        ("Processing", "Python + DuckDB", "التحميل، التحقق، الربط، والتجميع بكفاءة."),
        ("Analytical Storage", "Parquet + DuckDB", "حفظ بيانات مستخدمين معالجة وData Marts صغيرة."),
        ("Statistics", "SciPy + Statsmodels", "الاختبارات، فواصل الثقة، وأحجام الأثر."),
        ("Dashboard", "Streamlit + Plotly", "الفلاتر، الرسوم، الجداول، والتنزيل."),
        ("Quality", "Pytest + Data Checks", "اختبار الحسابات والجودة والانحدار."),
        ("Deployment", "GitHub + Streamlit Cloud", "نشر الكود والنتائج المجمعة فقط."),
    ], [1800, 2500, 5060], font_size=9.2)
    add_callout(doc, "تدفق البيانات", "Raw CSV → Data Audit → Clean/User-Level Features → Statistical Analysis → Aggregated Data Marts → Streamlit Dashboard → Report & Presentation", LIGHT_BLUE, BLUE)

    add_heading(doc, "9. خطة Data Pipeline", 1)
    phases = [
        ("9.1 الفهرسة والفحص", [
            "توثيق أسماء الملفات، الأحجام، Hash، الأعمدة، أنواع البيانات، ونطاق التواريخ.",
            "حساب التكرارات والقيم المفقودة والقيم غير الصالحة وتوزيع الدول والقنوات والـVariations.",
            "إنشاء تقرير Data Audit محفوظ داخل logs وdocs.",
        ]),
        ("9.2 التنظيف والتحويل", [
            "تحويل التواريخ والأرقام مع تسجيل الصفوف غير القابلة للتحويل.",
            "توحيد missing/Unknown دون حذفها بصمت.",
            "إزالة التكرارات المؤكدة فقط مع تسجيل العدد والسبب.",
            "تحويل البيانات الكبيرة إلى Parquet مضغوط ومقسّم عند الحاجة.",
        ]),
        ("9.3 بناء جدول المستخدم التحليلي", [
            "تجميع الإيرادات والسحوبات والأرباح وعدد المعاملات لكل user_id.",
            "Left join مع جميع المستخدمين لضمان بقاء المستخدمين دون معاملات.",
            "اشتقاق Net Revenue، monetized، has_cashout، first/last activity، وأيام النشاط.",
        ]),
        ("9.4 إنشاء Data Marts", [
            "ملخص عام، ملخص Variation، ملخص Channel، Cohort retention، Ad formats، Statistical tests.",
            "منع وجود user_id داخل ملفات data/marts القابلة للنشر.",
        ]),
    ]
    for title, bullets in phases:
        add_heading(doc, title, 2)
        for bullet in bullets:
            add_bullet(doc, bullet)

    add_heading(doc, "10. تعريف المؤشرات", 1)
    add_table(doc, ["المؤشر", "التعريف التنفيذي"], [
        ("Gross Revenue", "مجموع revenue_amount للمجموعة والفترة المحددتين."),
        ("Cashout", "مجموع cashout_amount المدفوع للمستخدمين."),
        ("Net Revenue", "Gross Revenue − Cashout."),
        ("ARPU", "Gross Revenue ÷ جميع المستخدمين المؤهلين في المجموعة، بمن فيهم أصحاب الإيراد الصفري."),
        ("Average Net Revenue/User", "Net Revenue ÷ جميع المستخدمين المؤهلين."),
        ("Activation Rate", "عدد activated=1 ÷ عدد المستخدمين المؤهلين."),
        ("Monetization Rate", "عدد المستخدمين ذوي total_revenue>0 ÷ عدد المستخدمين المؤهلين."),
        ("Exact-day Activity Dn", "وجود حدث مالي/إعلاني للمستخدم في اليوم n بالضبط بعد التثبيت."),
        ("Rolling Activity Dn+", "وجود حدث في اليوم n أو بعده؛ يعرض منفصلاً ولا يسمى Exact Retention."),
        ("Cohort Eligibility", "لا يدخل المستخدم في Dn إلا إذا كانت نهاية الرصد تسمح بمرور n يوماً كاملاً بعد تثبيته."),
        ("ROAS", "Revenue أو Net Revenue ÷ Ad Spend؛ لا يحسب قبل توفر تكلفة القناة أو الحملة."),
    ], [2800, 6560], font_size=8.9)
    add_callout(doc, "قيد منهجي مهم", "البيانات الحالية لا تحتوي جلسات دخول مستقلة؛ لذلك يجب تسمية المؤشر Activity Retention based on revenue-related events، وعدم عرضه على أنه Retention سلوكي كامل دون توضيح.", LIGHT_GOLD, GOLD)

    page_break(doc)
    add_heading(doc, "11. خطة التحليل الإحصائي", 1)
    for item in [
        "إحصاءات وصفية لكل Variation: العدد، المتوسط، الوسيط، الانحراف، الفواصل الربعية، والإيراد الصفري.",
        "Welch t-test لمقارنة المتوسطات عند عدم افتراض تساوي التباين.",
        "Mann–Whitney U كتحليل حساسية لتوزيعات الإيراد المنحرفة.",
        "Cohen’s d أو حجم أثر مناسب مع فاصل ثقة، وليس p-value وحدها.",
        "تصحيح تعدد المقارنات باستخدام Holm أو Benjamini–Hochberg حسب نوع الاستنتاج.",
        "مقارنات رئيسية محددة مسبقاً: كل Variation مقابل default، ثم المقارنات الزوجية الثانوية.",
        "تحليل نسب الاحتفاظ مع فواصل ثقة واختبار فروق النسب عند ملاءمته.",
        "عرض النتيجة بثلاثة مستويات: وصفياً أعلى، دال إحصائياً، وذو أثر عملي مفيد.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "12. متطلبات الداشبورد", 1)
    add_table(doc, ["الصفحة", "المحتوى الأساسي", "الفلاتر"], [
        ("Executive Overview", "KPIs، الإيراد، السحوبات، Net Revenue، ARPU، ومخطط زمني.", "الدولة، التاريخ، Variation، Channel"),
        ("A/B Variations", "الترتيب، الفواصل، الاختبارات، الأثر، والمقارنة مع default.", "المؤشر، الدولة، Cohort"),
        ("Retention & Cohorts", "D1/D3/D7/D14/D30 ومنحنيات Cohort مع أهلية الرصد.", "نوع retention، Cohort، Variation، Channel"),
        ("Acquisition Channels", "الحجم، ARPU، Net Revenue، retention، وleaderboard.", "الدولة، الحد الأدنى للحجم، الفترة"),
        ("Revenue & Ads", "Revenue type، Ad formats، cashout، ومزيج الإيرادات.", "Variation، Channel، Format"),
        ("Statistical Testing", "p-values المصححة، effect sizes، confidence intervals.", "المؤشر والمجموعتان"),
        ("Data Quality", "التغطية، المفقود، التكرار، التواريخ، وتعريفات KPIs.", "الملف ونوع المشكلة"),
        ("Export", "تنزيل CSV/PNG وملخص تنفيذي للفلتر الحالي.", "حسب الصفحة الحالية"),
    ], [1900, 4760, 2700], font_size=8.5)

    add_heading(doc, "13. المتطلبات الوظيفية وغير الوظيفية", 1)
    add_heading(doc, "13.1 وظيفية", 2)
    for item in [
        "تطبيق الفلاتر على جميع الرسوم والأرقام بصورة متسقة.",
        "إظهار تعريف المؤشر ومصدره بجانب كل نتيجة مهمة.",
        "تنزيل الجداول المفلترة والرسوم المختارة.",
        "تمييز النتائج غير الدالة إحصائياً وعدم تحويلها إلى توصيات قطعية.",
        "عرض وقت آخر تحديث وإصدار Pipeline الذي أنشأ ملفات النتائج.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "13.2 غير وظيفية", 2)
    for item in [
        "الأداء: تحميل الصفحات الأساسية سريعاً باستخدام caching وmarts صغيرة.",
        "الأمان: عدم تضمين بيانات شخصية أو معرفات مستخدمين في التطبيق المنشور.",
        "القابلية للصيانة: فصل منطق التحليل عن منطق العرض.",
        "القابلية لإعادة الإنتاج: تثبيت إصدارات المكتبات وإعدادات التحليل.",
        "سهولة الاستخدام: تصميم واضح ومتجاوب مع شرح مختصر للمؤشرات.",
    ]:
        add_bullet(doc, item)

    page_break(doc)
    add_heading(doc, "14. خطة التنفيذ الزمنية", 1)
    timeline_rows = [
        ("1", "تثبيت المتطلبات وفحص البيانات", "Data inventory، Hash، Schema، Quality report", "تقرير فحص وقرارات منهجية"),
        ("2", "بناء ingestion وParquet", "قراءة DuckDB، الأنواع، السجلات، تقسيم البيانات", "Processed layer قابلة لإعادة البناء"),
        ("3", "جدول المستخدم وKPIs", "Aggregation، join، validation، تعريفات المؤشرات", "User analytical table وKPI tests"),
        ("4", "Retention وCohorts", "الأهلية الزمنية، exact-day، rolling، cohorts", "Retention marts مدققة"),
        ("5", "التحليل الإحصائي", "Welch، Mann–Whitney، effect size، corrections", "Statistical results وinterpretation"),
        ("6", "قنوات الاستحواذ والإعلانات", "Leaderboards، countries، formats، revenue mix", "Channel/ad marts ورسوم"),
        ("7", "Dashboard - الجزء الأول", "Overview، Variations، Retention", "نسخة داخلية قابلة للاستخدام"),
        ("8", "Dashboard - الجزء الثاني", "Channels، Ads، Statistics، Quality، Export", "MVP مكتمل"),
        ("9", "الاختبار والتحسين والنشر", "Pytest، performance، privacy، GitHub، Streamlit", "Release candidate ورابط عرض"),
        ("10", "التقرير والعرض والتدريب", "الفصول، screenshots، slides، demo script", "حزمة التسليم النهائية"),
    ]
    add_table(doc, ["الأسبوع", "المحور", "الأعمال", "المخرج"], timeline_rows, [850, 2100, 3600, 2810], font_size=8.2)

    add_heading(doc, "15. خطة الاختبار وضمان الجودة", 1)
    add_table(doc, ["نوع الاختبار", "ما الذي يتحقق منه"], [
        ("Schema tests", "وجود الأعمدة المطلوبة وأنواعها ونطاقات القيم الأساسية."),
        ("Data reconciliation", "تطابق عدد المستخدمين ومجاميع الإيراد والسحب قبل وبعد التحويل."),
        ("Unit tests", "صحة معادلات Net Revenue وARPU وRetention والأهلية الزمنية."),
        ("Statistical tests", "نتائج ثابتة على بيانات اختبار معروفة وتصحيح المقارنات."),
        ("Integration tests", "تشغيل Pipeline كاملة وإنتاج جميع marts دون أخطاء."),
        ("Dashboard tests", "عمل الفلاتر والتنزيل وعدم تعارض الأرقام بين الصفحات."),
        ("Performance tests", "زمن التحميل والذاكرة وحجم الملفات المجمعة."),
        ("Privacy tests", "عدم وجود user_id أو سجلات فردية في ملفات النشر."),
        ("Regression checks", "التنبيه عند تغير الأرقام الرئيسية بعد تعديل الكود."),
    ], [2500, 6860], font_size=9.1)

    add_heading(doc, "16. خطة النشر", 1)
    add_numbered(doc, "تشغيل Pipeline الكاملة محلياً على النسخة الموجودة في data/raw.")
    add_numbered(doc, "مراجعة تقارير التحقق ثم اعتماد data/marts فقط للنشر.")
    add_numbered(doc, "رفع الكود والوثائق والملفات المجمعة الآمنة إلى GitHub خاص.")
    add_numbered(doc, "نشر Streamlit من المستودع الخاص وضبط Secrets إن وجدت.")
    add_numbered(doc, "تجهيز نسخة محلية تعمل دون إنترنت للمناقشة.")
    add_numbered(doc, "تسجيل رقم الإصدار وتاريخ البيانات وآخر تشغيل داخل الداشبورد.")

    add_heading(doc, "17. المخرجات النهائية", 1)
    for item in [
        "مستودع FYP II منظم مع README وتعليمات تشغيل واضحة.",
        "Data Pipeline قابلة لإعادة التشغيل على البيانات الكاملة.",
        "Parquet/DuckDB وData Marts مجمعة ومدققة.",
        "تحليل إحصائي كامل وتقرير نتائج قابل للتتبع.",
        "Dashboard Streamlit متعددة الصفحات.",
        "اختبارات آلية وتقارير جودة وأداء وخصوصية.",
        "نسخة محلية ونسخة سحابية للعرض.",
        "تقرير مشروع التخرج الثاني وعرض المناقشة ودليل Demo.",
    ]:
        add_bullet(doc, item)

    page_break(doc)
    add_heading(doc, "18. سجل المخاطر وخطط الاستجابة", 1)
    add_table(doc, ["الخطر", "الاحتمال/الأثر", "الاستجابة"], [
        ("حجم ملف الإيرادات يسبب بطئاً أو نفاد الذاكرة", "متوسط / مرتفع", "DuckDB وParquet ومعالجة دفعية وعدم تحميل الملف كاملاً في Pandas."),
        ("تعريف Retention غير دقيق", "مرتفع / مرتفع", "فصل exact-day وrolling، تطبيق أهلية الرصد، وتوضيح أنه مبني على أحداث الإيراد."),
        ("عدم وجود فروق دالة بين Variations", "متوسط / متوسط", "اعتبارها نتيجة صحيحة، والتركيز على الفواصل والأثر وعدم انتقاء النتائج."),
        ("غياب تكلفة القنوات", "مرتفع / متوسط", "عدم ادعاء ROAS؛ عرض Revenue/Retention وطلب ملف التكلفة كمدخل اختياري."),
        ("تسرب بيانات حساسة", "منخفض / مرتفع", "عدم رفع raw/processed، فحص ملفات النشر، واستخدام مستودع خاص."),
        ("انقطاع الإنترنت أثناء المناقشة", "متوسط / متوسط", "نسخة محلية كاملة ولقطات احتياطية."),
        ("اختلاف الأرقام بين التقرير والداشبورد", "متوسط / مرتفع", "مصدر موحد للمؤشرات واختبارات reconciliation ونسخة إصدار ثابتة."),
        ("تأخر الجدول", "متوسط / متوسط", "تثبيت MVP أولاً ونقل ML والتنبؤ إلى قائمة التوسع."),
    ], [2600, 1900, 4860], font_size=8.7)

    add_heading(doc, "19. خطة التقرير النهائي والعرض", 1)
    add_heading(doc, "19.1 فصول التقرير المقترحة", 2)
    for item in [
        "مقدمة ومشكلة المشروع وأهداف FYP II.",
        "مراجعة العمل السابق والانتقال من عينة FYP I إلى البيانات الكاملة.",
        "وصف البيانات وحوكمتها وجودتها.",
        "تصميم النظام والمعمارية وData Pipeline.",
        "تعريف المؤشرات ومنهجية Cohorts والاختبارات الإحصائية.",
        "النتائج والمناقشة والتوصيات التجارية.",
        "تصميم الداشبورد والتنفيذ والاختبار والنشر.",
        "القيود والعمل المستقبلي والخاتمة.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "19.2 سيناريو المناقشة", 2)
    for item in [
        "دقيقتان: المشكلة وقيمة المشروع.",
        "ثلاث دقائق: البيانات والمعمارية وتحسين FYP II.",
        "خمس دقائق: Demo للداشبورد ومسار قرار نموذجي.",
        "ثلاث دقائق: النتائج الإحصائية وما تعنيه عملياً.",
        "دقيقتان: الاختبارات والخصوصية والقيود والعمل المستقبلي.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "20. قواعد إدارة المشروع", 1)
    for item in [
        "لا يُكتب منطق نهائي داخل Notebook فقط؛ ينقل إلى وحدات Python قابلة للاختبار.",
        "كل تغيير في تعريف KPI يوثق في docs ويحدث في الكود والداشبورد معاً.",
        "تُحفظ البيانات الخام للقراءة فقط، ولا تُعدل يدوياً.",
        "تُستخدم فروع Git واضحة ورسائل Commit تصف التغيير.",
        "تُثبت إصدارات المكتبات في requirements.txt أو pyproject.toml.",
        "تُحفظ نتائج كل تشغيل مع وقت التشغيل وإصدار الكود وتقرير التحقق.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "21. تعريف الاكتمال النهائي", 1)
    add_callout(doc, "يُعد المشروع مكتملاً عندما", "تعمل Pipeline على جميع البيانات، وتنجح اختبارات المصالحة والخصوصية، وتنتج Data Marts معتمدة، ويعرض الداشبورد النتائج والفلاتر والتنزيل بصورة صحيحة، وتعمل النسختان المحلية والسحابية، ويتطابق التقرير والعرض مع الإصدار النهائي للأرقام.", LIGHT_TEAL, TEAL)

    add_heading(doc, "22. أول سبع خطوات تنفيذية", 1)
    for step in [
        "إنشاء ملف manifest وSHA-256 للنسخ الخام.",
        "إعداد README و.gitignore ومنع رفع data/raw وdata/processed.",
        "إنشاء Data Audit script وتقرير جودة أولي.",
        "تثبيت قاموس البيانات وتعريفات KPIs وRetention.",
        "تحويل البيانات إلى Parquet وبناء DuckDB محلي.",
        "إنشاء جدول المستخدم التحليلي واختبارات المصالحة.",
        "بناء أول Data Mart وصفحة Executive Overview كنموذج أولي.",
    ]:
        add_numbered(doc, step)

    add_paragraph(doc, "نهاية الخطة — الإصدار 1.0", size=9, bold=True, color=MUTED,
                  before=20, align=WD_ALIGN_PARAGRAPH.CENTER)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
